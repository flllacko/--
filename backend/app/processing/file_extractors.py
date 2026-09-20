from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from .web import WebExtractionError, extract_web_page


class FileExtractionError(ValueError):
    """文件不能可靠提取为正文。"""


MAX_FILE_BYTES = 50 * 1024 * 1024


@dataclass(frozen=True)
class ExtractedDocument:
    path: Path
    title: str
    text: str
    media_type: str


def extract_file(file_path: Path, *, max_bytes: int = MAX_FILE_BYTES) -> ExtractedDocument:
    path = file_path.resolve()
    if not path.is_file():
        raise FileExtractionError(f"文件不存在：{path}")
    if path.stat().st_size > max_bytes:
        raise FileExtractionError(f"文件超过 50 MB 限制：{path.name}")

    extension = path.suffix.lower()
    if extension == ".txt":
        return _from_text_file(path)
    if extension == ".md":
        return _from_text_file(path, media_type="Markdown")
    if extension == ".pdf":
        return _from_pdf(path)
    if extension == ".docx":
        return _from_docx(path)
    if extension == ".doc":
        raise FileExtractionError("旧式 .doc 暂不支持，请转换为 .docx 或 PDF 后重试")
    if extension in {".xlsx", ".xls"}:
        return _from_spreadsheet(path, extension)
    if extension in {".html", ".htm"}:
        return _from_html(path)
    raise FileExtractionError(f"暂不支持的文件类型：{extension or '无扩展名'}")


def clean_extracted_text(text: str) -> str:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    lines: list[str] = []
    for raw_line in normalized.splitlines():
        line = re.sub(r"[ \t]+", " ", raw_line).strip()
        if line:
            lines.append(line)
        elif lines and lines[-1] != "":
            lines.append("")
    while lines and lines[-1] == "":
        lines.pop()
    return "\n".join(lines).strip()


def _from_text_file(path: Path, media_type: str = "纯文本") -> ExtractedDocument:
    raw = path.read_bytes()
    text: str | None = None
    for encoding in ("utf-8-sig", "gb18030", "utf-16"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise FileExtractionError(f"无法识别文本编码：{path.name}")
    cleaned = clean_extracted_text(text)
    if not cleaned:
        raise FileExtractionError(f"正文为空：{path.name}")
    title_match = re.search(r"^#\s+(.+?)\s*$", cleaned, re.MULTILINE)
    title = title_match.group(1).strip() if title_match else path.stem
    return ExtractedDocument(path=path, title=title, text=cleaned, media_type=media_type)


def _from_pdf(path: Path) -> ExtractedDocument:
    from pypdf import PdfReader

    try:
        reader = PdfReader(str(path))
        page_parts: list[str] = []
        for index, page in enumerate(reader.pages, start=1):
            text = clean_extracted_text(page.extract_text() or "")
            if text:
                page_parts.append(f"## 第 {index} 页\n\n{text}")
    except Exception as exc:  # noqa: BLE001 - 第三方解析器错误转为统一失败
        raise FileExtractionError(f"PDF 解析失败：{path.name}：{exc}") from exc
    content = clean_extracted_text("\n\n".join(page_parts))
    if not content:
        raise FileExtractionError("PDF 没有可提取文字，可能是扫描件；当前版本不启用 OCR")
    return ExtractedDocument(path=path, title=path.stem, text=content, media_type="PDF")


def _from_docx(path: Path) -> ExtractedDocument:
    from docx import Document

    try:
        document = Document(str(path))
        parts: list[str] = []
        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                parts.append(paragraph.text)
        for table_index, table in enumerate(document.tables, start=1):
            parts.append(f"## 表格 {table_index}")
            for row in table.rows:
                values = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                if any(values):
                    parts.append(" | ".join(values))
    except Exception as exc:  # noqa: BLE001
        raise FileExtractionError(f"DOCX 解析失败：{path.name}：{exc}") from exc
    content = clean_extracted_text("\n\n".join(parts))
    if not content:
        raise FileExtractionError(f"DOCX 正文为空：{path.name}")
    title = next((line[2:].strip() for line in content.splitlines() if line.startswith("# ")), path.stem)
    return ExtractedDocument(path=path, title=title, text=content, media_type="DOCX")


def _from_spreadsheet(path: Path, extension: str) -> ExtractedDocument:
    try:
        if extension == ".xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
            sheets = []
            for worksheet in workbook.worksheets:
                rows = []
                for row in worksheet.iter_rows(values_only=True):
                    values = ["" if value is None else str(value).strip() for value in row]
                    if any(values):
                        rows.append(" | ".join(values))
                if rows:
                    sheets.append(f"## 工作表：{worksheet.title}\n\n" + "\n".join(rows))
            workbook.close()
        else:
            import xlrd

            workbook = xlrd.open_workbook(str(path), on_demand=True)
            sheets = []
            for sheet in workbook.sheets():
                rows = []
                for row_index in range(sheet.nrows):
                    values = [str(value).strip() for value in sheet.row_values(row_index)]
                    if any(values):
                        rows.append(" | ".join(values))
                if rows:
                    sheets.append(f"## 工作表：{sheet.name}\n\n" + "\n".join(rows))
            workbook.release_resources()
    except Exception as exc:  # noqa: BLE001
        raise FileExtractionError(f"表格解析失败：{path.name}：{exc}") from exc
    content = clean_extracted_text("\n\n".join(sheets))
    if not content:
        raise FileExtractionError(f"表格没有可提取内容：{path.name}")
    return ExtractedDocument(path=path, title=path.stem, text=content, media_type="Excel")


def _from_html(path: Path) -> ExtractedDocument:
    raw = path.read_bytes()
    html = raw.decode("utf-8", errors="replace")
    try:
        page = extract_web_page(html, f"https://local.invalid/{path.name}")
    except WebExtractionError as exc:
        raise FileExtractionError(str(exc)) from exc
    return ExtractedDocument(path=path, title=page.title, text=page.body_text, media_type="HTML")

