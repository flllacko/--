from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from app.processing.file_extractors import FileExtractionError, extract_file
from app.processing.lance_store import has_document
from app.processing.processor import process_extracted_file


class FileExtractorsTest(unittest.TestCase):
    def test_extracted_file_uses_the_real_vector_ingest_boundary(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory)
            root = base / "资料原件"
            root.joinpath("学生资助").mkdir(parents=True)
            source = root / "学生资助" / "policy.txt"
            source.write_text("国家助学金按照公开政策分为三档，具体以学校通知为准。", encoding="utf-8")
            database = base / "app.db"
            index = base / "index"

            with patch("app.processing.processor.embed_texts", return_value=[[0.1, 0.2, 0.3]]):
                result = process_extracted_file(
                    source,
                    source_root=root,
                    db_path=database,
                    index_path=index,
                )

            self.assertEqual(result.status, "处理成功")
            self.assertTrue(has_document(index, int(result.document_id)))

    def test_text_docx_and_xlsx_are_extracted(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory)
            text_path = base / "notice.txt"
            text_path.write_bytes("标题\n\n国家助学金三档".encode("utf-8"))
            self.assertIn("国家助学金三档", extract_file(text_path).text)

            from docx import Document

            docx_path = base / "notice.docx"
            document = Document()
            document.add_heading("DOCX 通知", level=1)
            document.add_paragraph("DOCX 正文测试")
            document.save(docx_path)
            docx_result = extract_file(docx_path)
            self.assertEqual(docx_result.media_type, "DOCX")
            self.assertIn("DOCX 正文测试", docx_result.text)

            from openpyxl import Workbook

            xlsx_path = base / "notice.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "资助标准"
            worksheet.append(["项目", "金额"])
            worksheet.append(["国家助学金", "4900元"])
            workbook.save(xlsx_path)
            workbook.close()
            xlsx_result = extract_file(xlsx_path)
            self.assertEqual(xlsx_result.media_type, "Excel")
            self.assertIn("国家助学金 | 4900元", xlsx_result.text)

    def test_html_and_unsupported_doc_are_explicit(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory)
            html_path = base / "notice.html"
            html_path.write_text(
                "<main><h1>网页通知</h1><p>这是网页正文，包含足够的通知内容，用于验证本地 HTML 能够经过正文提取后进入后续处理流程。</p></main>",
                encoding="utf-8",
            )
            html_result = extract_file(html_path)
            self.assertEqual(html_result.media_type, "HTML")
            self.assertIn("网页正文", html_result.text)

            doc_path = base / "legacy.doc"
            doc_path.write_bytes(b"not a doc")
            with self.assertRaises(FileExtractionError):
                extract_file(doc_path)


if __name__ == "__main__":
    unittest.main()
